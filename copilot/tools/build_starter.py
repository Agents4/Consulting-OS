# Builds copilot/client-workspace-starter/ and Client-Workspace-Starter.zip
# Needs pandoc, python-docx, openpyxl. Run: python3 copilot/tools/build_starter.py
import os, re, shutil, subprocess, zipfile, glob
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
KIT = ROOT + "/gemini/client-kit"
OUT = ROOT + "/copilot/client-workspace-starter"
TOP = OUT + "/CLIENT-NAME"
if os.path.exists(OUT): shutil.rmtree(OUT)
for d in ["00_Context", "01_Outputs", "02_Data", "03_Trackers", "04_Archive", "05_Templates"]:
    os.makedirs(f"{TOP}/{d}")

subs = ["Proposals and Workplans", "Hypotheses and Issue Trees", "Interviews and Research", "Analyses and Models",
        "Storyline and Decks", "Reviews", "SteerCo and Status", "Meetings and Decisions",
        "Recommendations and Roadmaps", "Transition and Retro"]
for n in subs:
    os.makedirs(f"{TOP}/01_Outputs/{n}")
    open(f"{TOP}/01_Outputs/{n}/_folder-note.txt", "w").write(f"{n}: final and working files for this topic. Name files Topic_YYYY-MM-DD_vN.\n")
open(f"{TOP}/02_Data/_folder-note.txt", "w").write("Client data extracts. Store only what the client approved for this platform. Record each file in doc 04, Data received.\n")
open(f"{TOP}/04_Archive/_folder-note.txt", "w").write("Old versions and retired context. Move log entries older than 6 weeks here.\n")

# context docs
swap = [
 ("- Models and trackers: Google Sheets", "- Models and trackers: Excel"),
 ("- Decks: Google Slides", "- Decks: PowerPoint"),
 ("- Memos: Google Docs", "- Memos: Word"),
 ("[Client data room / Firm Drive / Google Workspace (approval recorded at: [where])]", "[Client data room / Firm SharePoint / Microsoft 365 tenant (approval recorded at: [where])]"),
]
tmp = ROOT + "/copilot/_tmp"; os.makedirs(tmp, exist_ok=True)
for f in sorted(glob.glob(KIT + "/*.md")):
    t = open(f, encoding="utf-8").read()
    for a, b in swap: t = t.replace(a, b)
    name = os.path.basename(f)[:-3]
    md = f"{tmp}/{name}.md"; open(md, "w", encoding="utf-8").write(t)
    out = f"{TOP}/00_Context/{name}.docx"
    subprocess.run(["pandoc", md, "-f", "markdown-smart", "-o", out], check=True)
    d = docx.Document(out)
    for tb in d.tables:
        tblPr = tb._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement("w:" + edge)
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "808080")
            borders.append(e)
        tblPr.append(borders)
        for c in tb.rows[0].cells:
            for p in c.paragraphs:
                for r in p.runs: r.font.bold = True
    d.save(out)
shutil.rmtree(tmp)

# trackers workbook
wb = Workbook()
hdr_font = Font(name="Arial", bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor="1F3864")
body = Font(name="Arial", size=10)
def sheet(ws, headers, rows, widths):
    ws.append(headers)
    for r in rows: ws.append(r)
    for c in ws[1]:
        c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths, 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = body; c.alignment = Alignment(wrap_text=True, vertical="top")
ws = wb.active; ws.title = "Data Requests"
sheet(ws, ["ID","Dataset","Fields and grain","Period","Purpose (hypothesis or slide)","Client owner","Requested","Needed by","Received","Format","Quality","Location","Status"],
      [["D1","","","","","","","","","","","","Draft"]], [6,22,28,12,26,16,12,12,12,10,12,22,12])
dv = DataValidation(type="list", formula1='"Draft,Requested,Promised,Partial,Received,Cleaned,Blocked,Not available"', allow_blank=True)
ws.add_data_validation(dv); dv.add("M2:M200")
ws2 = wb.create_sheet("Risk Register")
sheet(ws2, ["ID","Risk","Category (Data / Scope / People / Client / Quality / Confidentiality)","Likelihood (1-5)","Impact (1-5)","Score","Early signal","Mitigation","Owner","Status"],
      [["R1","","","","",'=IF(AND(ISNUMBER(D2),ISNUMBER(E2)),D2*E2,"")',"","","","Open"]], [6,34,26,12,10,8,24,28,14,10])
for r in range(3, 101):
    ws2.cell(r, 6).value = f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),D{r}*E{r},"")'
    for c in range(1, 11): ws2.cell(r, c).font = body
dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
ws2.add_data_validation(dv2); dv2.add("D2:E100")
ws3 = wb.create_sheet("Workplan")
sheet(ws3, ["Workstream","Task","Owner","Start","End","Status","Depends on","Gate","Hypothesis"],
      [["","","","","","Not started","","",""]], [18,34,14,12,12,14,16,14,14])
dv3 = DataValidation(type="list", formula1='"Not started,In progress,Blocked,Done"', allow_blank=True)
ws3.add_data_validation(dv3); dv3.add("F2:F200")
wb.save(f"{TOP}/03_Trackers/Engagement-Trackers.xlsx")
# optional: recalculate so formula values are cached (needs the xlsx skill's recalc.py and LibreOffice)
for rc in glob.glob("/root/.claude/skills/synced/*/xlsx/scripts/recalc.py"):
    subprocess.run(["python3", rc, f"{TOP}/03_Trackers/Engagement-Trackers.xlsx"], check=False)
    break

for f in ["templates/model-starter-template.xlsx", "templates/ghost-deck-starter.pptx"]:
    shutil.copy(ROOT + "/" + f, f"{TOP}/05_Templates/")

# zip
zp = ROOT + "/copilot/Client-Workspace-Starter.zip"
if os.path.exists(zp): os.remove(zp)
with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
    for dp, dn, fn in os.walk(OUT):
        for f in fn:
            full = os.path.join(dp, f); z.write(full, os.path.relpath(full, OUT))
print("ok", os.path.getsize(zp)//1024, "KB")
